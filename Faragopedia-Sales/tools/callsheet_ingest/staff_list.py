"""Resolve whether a name on a callsheet is Farago staff or a freelancer.

Simon supplied the staff list on 2026-08-25 precisely because callsheets do not record
employment status. An earlier assumption that the "PRODUCTION | FARAGO PROJECTS" heading
identified staff was wrong: on the McQueen 23 Jan sheet only 1 of the 7 people under that
heading is on the list — the Producer and Production Manager are both freelance. The
heading gives the role grouping; this file gives the employment status.

The list is client HR data (names, job titles, start dates), so it is read from outside
the repo and never copied in. Point FARAGO_STAFF_LIST at it in .env.

Matching is first-given-name + surname, deliberately loose on middle names: the list
carries full legal names ("Robin Jane Elizabeth Ashby-Carter") while callsheets carry
working names ("Robin Ashby-Carter"). Strict matching silently marks staff as freelance.

Usage:
    from staff_list import StaffList
    staff = StaffList.load()
    staff.lookup("Robin Ashby-Carter")   # -> StaffMember(office="UK", ...) or None
"""

from __future__ import annotations

import os
import re
from dataclasses import dataclass
from pathlib import Path

DEFAULT_ENV_VAR = "FARAGO_STAFF_LIST"


def _tokens(value) -> list[str]:
    return re.sub(r"[^a-z]+", " ", str(value or "").lower()).split()


@dataclass(frozen=True)
class StaffMember:
    given: tuple[str, ...]
    surname: str
    title: str
    office: str
    start_date: str | None


class StaffList:
    def __init__(self, members: list[StaffMember]):
        self.members = members

    @classmethod
    def load(cls, path: str | Path | None = None) -> "StaffList":
        path = path or os.environ.get(DEFAULT_ENV_VAR)
        if not path:
            raise RuntimeError(
                f"{DEFAULT_ENV_VAR} is not set. Point it at the Farago staff list workbook; "
                "it is client HR data and is deliberately not kept in this repo."
            )
        from openpyxl import load_workbook

        wb = load_workbook(path, data_only=True)
        members: list[StaffMember] = []
        for ws in wb.worksheets:
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                continue
            header = [str(c or "").lower() for c in rows[0]]
            # The Paris tab puts surname first; UK and US put first name first.
            first_idx = 0 if header and "first" in header[0] else 1
            last_idx = 1 - first_idx
            title_idx = next(
                (i for i, h in enumerate(header) if "title" in h or "position" in h), 2
            )
            date_idx = next(
                (i for i, h in enumerate(header) if "date" in h), None
            )
            for row in rows[1:]:
                if not row or not any(row):
                    continue
                given = tuple(_tokens(row[first_idx] if first_idx < len(row) else ""))
                surname = _tokens(row[last_idx] if last_idx < len(row) else "")
                if not given or not surname:
                    continue
                start = None
                if date_idx is not None and date_idx < len(row) and row[date_idx]:
                    start = str(row[date_idx])[:10]
                members.append(
                    StaffMember(
                        given=given,
                        surname=surname[-1],
                        title=str(row[title_idx] or "").strip() if title_idx < len(row) else "",
                        office=ws.title,
                        start_date=start,
                    )
                )
        wb.close()
        return cls(members)

    def lookup(self, name: str) -> StaffMember | None:
        """Match on surname plus any given name. Returns None for freelancers.

        None means "not on the list", which is how Simon defined freelance — so an
        incomplete or stale list quietly reclassifies staff as freelance. Re-request the
        list rather than loosening this further.
        """
        parts = _tokens(name)
        if len(parts) < 2:
            return None
        given, surname = parts[0], parts[-1]
        for m in self.members:
            if surname == m.surname and given in m.given:
                return m
        return None

    def lookup_as_at(self, name: str, shoot_date: str | None) -> StaffMember | None:
        """Staff status as at a shoot date, not as at today.

        The list is an Aug-2026 snapshot against a 2025-2026 corpus, so someone who
        joined in 2026 was a freelancer on a 2025 sheet, and a plain lookup() would
        backdate their employment across the whole archive. This is the reason the
        decisions doc requires the list to be applied at ingestion and never frozen
        into callsheets-md/.

        A member with no start_date is treated as staff on any date: absence of a date
        is missing data, not evidence they had not joined. `unresolved_start_dates`
        reports how many such rows the list carries so it stays visible.
        """
        member = self.lookup(name)
        if member is None or not shoot_date or not member.start_date:
            return member
        return None if shoot_date < member.start_date else member

    @property
    def unresolved_start_dates(self) -> list[str]:
        return [f"{' '.join(m.given)} {m.surname}" for m in self.members if not m.start_date]

    def classify(self, name: str, shoot_date: str | None = None) -> str:
        member = self.lookup_as_at(name, shoot_date) if shoot_date else self.lookup(name)
        return f"staff ({member.office})" if member else "freelance"


if __name__ == "__main__":
    from dotenv import load_dotenv

    load_dotenv(Path(__file__).parent / ".env")
    staff = StaffList.load()
    print(f"{len(staff.members)} staff loaded")
    by_office: dict[str, int] = {}
    for m in staff.members:
        by_office[m.office] = by_office.get(m.office, 0) + 1
    for office, n in sorted(by_office.items()):
        print(f"  {office}: {n}")
